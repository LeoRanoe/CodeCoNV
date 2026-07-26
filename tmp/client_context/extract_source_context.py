from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def safe_text(value: object | None) -> str:
    if value is None:
        return ""
    return str(value).replace("\x00", "").strip()


def extract_pdf(path: Path) -> dict[str, object]:
    reader = PdfReader(path)
    pages: list[dict[str, object]] = []
    for index, page in enumerate(reader.pages, start=1):
        text = safe_text(page.extract_text())
        pages.append(
            {
                "page": index,
                "characters": len(text),
                "text": text,
            }
        )
    return {
        "kind": "pdf",
        "path": str(path),
        "pages": pages,
        "metadata": {str(k): safe_text(v) for k, v in (reader.metadata or {}).items()},
    }


def extract_docx(path: Path) -> dict[str, object]:
    document = Document(path)
    paragraphs = [
        {
            "index": index,
            "style": safe_text(paragraph.style.name if paragraph.style else ""),
            "text": safe_text(paragraph.text),
        }
        for index, paragraph in enumerate(document.paragraphs, start=1)
        if safe_text(paragraph.text)
    ]
    tables: list[dict[str, object]] = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = [
            [safe_text(cell.text) for cell in row.cells]
            for row in table.rows
        ]
        tables.append(
            {
                "table": table_index,
                "rows": len(rows),
                "columns": max((len(row) for row in rows), default=0),
                "values": rows,
            }
        )

    embedded_media: list[dict[str, object]] = []
    with zipfile.ZipFile(path) as archive:
        for info in archive.infolist():
            if info.filename.startswith("word/media/"):
                embedded_media.append(
                    {
                        "name": info.filename,
                        "bytes": info.file_size,
                    }
                )
        package_members = set(archive.namelist())

    properties = document.core_properties
    return {
        "kind": "docx",
        "path": str(path),
        "paragraphs": paragraphs,
        "tables": tables,
        "sections": len(document.sections),
        "inline_shapes": len(document.inline_shapes),
        "embedded_media": embedded_media,
        "has_comments_part": "word/comments.xml" in package_members,
        "has_footnotes_part": "word/footnotes.xml" in package_members,
        "has_endnotes_part": "word/endnotes.xml" in package_members,
        "properties": {
            "title": safe_text(properties.title),
            "subject": safe_text(properties.subject),
            "author": safe_text(properties.author),
            "last_modified_by": safe_text(properties.last_modified_by),
            "created": safe_text(properties.created),
            "modified": safe_text(properties.modified),
        },
    }


def main() -> int:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: extract_source_context.py <source_dir> <output_json>")

    source_dir = Path(sys.argv[1]).resolve()
    output_json = Path(sys.argv[2]).resolve()
    records: list[dict[str, object]] = []

    for path in sorted(source_dir.iterdir(), key=lambda item: item.name.casefold()):
        suffix = path.suffix.casefold()
        if suffix == ".pdf":
            records.append(extract_pdf(path))
        elif suffix == ".docx":
            records.append(extract_docx(path))

    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(
        json.dumps(records, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(
        json.dumps(
            [
                {
                    "file": Path(record["path"]).name,
                    "kind": record["kind"],
                    "pages": len(record.get("pages", [])),
                    "paragraphs": len(record.get("paragraphs", [])),
                    "tables": len(record.get("tables", [])),
                }
                for record in records
            ],
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
